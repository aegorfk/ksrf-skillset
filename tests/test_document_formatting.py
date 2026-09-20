"""Synthetic checks for the user-selected document formatting standard."""
import sys
import re
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/ksrf-complaint-cycle/lib'))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ksrf.filing.renderer import render_docx, _configure_document, _display_text, normalize_text
from ksrf.filing.working_draft import prepare_working_draft


class DocumentFormattingTests(unittest.TestCase):
    def complaint(self):
        sections = [
            ('addressee', 'Синтетический адресат'),
            ('applicant', 'Адрес: [УКАЗАТЬ АДРЕС]'),
            ('facts', 'Суд отказал в удовлетворении заявления по акту № 17.'),
            ('authorities', 'Условный источник №18: https://example.invalid/№19'),
            ('extra_reasoning', 'Содержательный дополнительный довод.'),
            ('enclosures', '1. Копия решения № 17.\n2. Извлечение из нормы.'),
        ]
        # One paragraph per numbered entry, as expected by the structured model.
        payload = {'matter_id':'synthetic', 'draft_id':'layout', 'sections':[
            {'code':code, 'heading':'Дополнительное основание', 'sentences':[
                {'text':line, 'role':'narrative', 'support_status':'pending'}
                for line in text.splitlines()]} for code,text in sections]}
        return prepare_working_draft(payload)[1]

    def render(self, folder):
        complaint = self.complaint()
        path = Path(folder) / 'complaint.docx'
        render_docx(complaint, path)
        return complaint, Document(path)

    def test_geometry_and_paragraph_style_values(self):
        with tempfile.TemporaryDirectory() as folder:
            _, document = self.render(folder)
            section = document.sections[0]
            for field, expected in {'page_width':210, 'page_height':297, 'left_margin':25,
                'right_margin':20, 'top_margin':20, 'bottom_margin':20,
                'header_distance':9, 'footer_distance':9}.items():
                self.assertAlmostEqual(getattr(section, field).mm, expected, places=1)
            expected = {
                'Normal': (12, 1.15, 0, 6, 10), 'Title':(14,1,12,6,0),
                'Subtitle':(12,1,0,12,0), 'KSRF Heading':(13,1.1,12,6,0),
                'KSRF Subheading':(12,1.1,10,6,0), 'KSRF Complaint Header':(11.5,1,0,4,0),
                'KSRF Source':(11,1,0,4,0), 'KSRF Footer':(11,1,0,0,0),
            }
            for name,(size,line,before,after,first) in expected.items():
                style=document.styles[name]; fmt=style.paragraph_format
                with self.subTest(style=name):
                    self.assertEqual(style.font.size.pt, size)
                    self.assertAlmostEqual(fmt.line_spacing, line)
                    self.assertEqual(fmt.space_before.pt, before)
                    self.assertEqual(fmt.space_after.pt, after)
                    self.assertAlmostEqual(fmt.first_line_indent.mm, first, places=1)
                    self.assertFalse(fmt.page_break_before)
            header=document.styles['KSRF Complaint Header'].paragraph_format
            self.assertAlmostEqual(header.left_indent.mm,82.5,places=1)
            self.assertEqual(header.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertEqual(document.styles['Normal'].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_fonts_language_complex_size_and_style_assignments(self):
        with tempfile.TemporaryDirectory() as folder:
            _, document=self.render(folder)
            used={p.style.name for p in document.paragraphs}|{'KSRF Footer'}
            for name in used:
                style=document.styles[name];rpr=style._element.rPr
                with self.subTest(style=name):
                    self.assertEqual(style.font.name,'Times New Roman')
                    self.assertEqual(str(style.font.color.rgb),'000000')
                    self.assertEqual(rpr.find(qn('w:szCs')).get(qn('w:val')),str(int(style.font.size.pt*2)))
                    self.assertEqual(rpr.find(qn('w:lang')).get(qn('w:val')),'ru-RU')
                    for attr in ('ascii','hAnsi','eastAsia','cs'):
                        self.assertEqual(rpr.rFonts.get(qn('w:'+attr)),'Times New Roman')
                    self.assertFalse(any(k.lower().endswith('theme') for k in rpr.rFonts.attrib))
                    self.assertEqual(rpr.find(qn('w:iCs')).get(qn('w:val')),'0')
                    self.assertIsNone(rpr.find(qn('w:spacing')))
                    self.assertIsNone(style._element.pPr.find(qn('w:pBdr')))
                    self.assertIsNone(style._element.pPr.find(qn('w:numPr')))
            paras={p.text:p for p in document.paragraphs}
            self.assertEqual(paras['ЖАЛОБА'].style.name,'Title')
            self.assertEqual(paras['на нарушение конституционных прав и свобод'].style.name,'Subtitle')
            self.assertEqual(paras['Дополнительное основание'].style.name,'KSRF Subheading')
            source=next(p for p in document.paragraphs if p.text.startswith('Условный источник'))
            self.assertEqual(source.style.name,'KSRF Source')
            self.assertEqual(document.sections[0].footer.paragraphs[0].style.name,'KSRF Footer')
            self.assertTrue(document.styles['Title'].font.bold)
            self.assertFalse(document.styles['Subtitle'].font.bold)
            for p in document.paragraphs:
                self.assertIsNone(p.paragraph_format.alignment)
                for run in p.runs:
                    self.assertIsNone(run.font.size)
                    self.assertIsNone(run.font.name)

    def test_source_text_numbers_links_and_highlights_survive(self):
        complaint=self.complaint()
        before=complaint.to_dict()
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'complaint.docx';render_docx(complaint,path)
            document=Document(path);text='\n'.join(p.text for p in document.paragraphs)
            self.assertIn('Суд отказал',text)
            self.assertIn('№\u00a017',text)
            self.assertIn('https://example.invalid/№19',text)
            self.assertIn('1.\tКопия решения №\u00a017.',text)
            self.assertIn('2.\tИзвлечение из нормы.',text)
            for section in complaint.sections:
                for sentence in section.sentences:
                    self.assertIn(re.sub(r'\s','',sentence.text),re.sub(r'\s','',text))
            self.assertEqual(complaint.to_dict(),before)
            highlighted=[r.text for p in document.paragraphs for r in p.runs
                         if r.font.highlight_color==WD_COLOR_INDEX.YELLOW]
            self.assertIn('[УКАЗАТЬ АДРЕС]',highlighted)
            self.assertFalse(document._element.xpath('.//w:br[@w:type="page"]'))
            self.assertFalse(document._element.xpath('.//w:numPr'))
            numbered=[p for p in document.paragraphs if p.style.name=='KSRF Numbered Item']
            self.assertEqual(len(numbered),2)
            fmt=document.styles['KSRF Numbered Item'].paragraph_format
            self.assertAlmostEqual(fmt.left_indent.mm,6.5,places=1)
            self.assertAlmostEqual(fmt.first_line_indent.mm,-6.5,places=1)
            self.assertEqual(fmt.alignment,WD_ALIGN_PARAGRAPH.LEFT)

    def test_style_setup_preserves_existing_fields_and_links(self):
        document=Document();p=document.add_paragraph('Исходный текст ')
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'DATE');p._p.append(field)
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'section1');p._p.append(link)
        before=p._p.xml
        _configure_document(document,self.complaint())
        self.assertEqual(p._p.xml,before)
        self.assertEqual(_display_text('№ 1; №2; https://example.invalid/№3'),'№\u00a01; №\u00a02; https://example.invalid/№3')


if __name__=='__main__':
    unittest.main()
