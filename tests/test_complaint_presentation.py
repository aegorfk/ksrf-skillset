"""Synthetic regressions for court document / review separation."""
import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

LIB = Path(__file__).resolve().parents[1] / "skills/ksrf-complaint-cycle/lib"
sys.path.insert(0, str(LIB))

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from ksrf.filing.presentation import COMPLAINT_TITLE, FACTS_HEADING, REASONING_HEADING, REQUEST_HEADING, PRAYER_MARKER, PresentationError, complaint_blocks
from ksrf.filing.renderer import render_docx, render_review_markdown, complaint_plain_text, find_unresolved_placeholders
from ksrf.filing.working_draft import prepare_working_draft, _review_markdown


class ComplaintPresentationTests(unittest.TestCase):
    def make_complaint(self):
        def section(code, text, role='narrative'):
            return {'code':code, 'heading':'Internal '+code,
                    'sentences':[{'text':text, 'role':role, 'support_status':'pending'}]}
        return prepare_working_draft({'matter_id':'synthetic-matter', 'draft_id':'synthetic-draft',
            'title':'Internal draft title', 'sections':[
                section('addressee','Синтетический адресат'),
                section('applicant','Адрес: [УКАЗАТЬ АДРЕС]'),
                section('facts','Суд отказал в удовлетворении заявления.', 'court_reasoning'),
                section('rights_analysis','Оспариваемый смысл не соответствует Конституции.', 'fact'),
                section('adverse_material',
                        'Суд установил, что отсутствие необходимых материалов не является доказанным; эта позиция требует разграничения.',
                        'court_reasoning'),
                section('review_notes','Риск: требуется проверить опоры довода.'),
            ]})

    def test_court_document_keeps_legal_negation_and_adverse_finding(self):
        original, complaint, gaps = self.make_complaint()
        text = complaint_plain_text(complaint)
        self.assertIn('Суд отказал', text)
        self.assertIn('не соответствует Конституции', text)
        self.assertIn('отсутствие необходимых материалов не является доказанным', text)
        for absent in ('РАБОЧИЙ ПРОЕКТ','ПРОВЕРИТЬ:', 'synthetic-matter', 'Internal draft title',
                       'Риск:','Возможные возражения','Представитель'):
            self.assertNotIn(absent, text)
        report = _review_markdown(original, gaps)
        self.assertIn('Риск: требуется проверить', report)
        self.assertIn('sent-', report)
        self.assertEqual(complaint.sections[2].sentences[0].support_status, 'pending')
        self.assertEqual(complaint.sections[2].sentences[0].text,
                         original.sections[2].sentences[0].text)

    def test_all_substantive_source_sentences_survive_presentation(self):
        original, complaint, _ = self.make_complaint()
        text = complaint_plain_text(complaint)
        for section in original.sections:
            for sentence in section.sentences:
                if section.code == 'review_notes':
                    self.assertNotIn(sentence.text, text)
                else:
                    self.assertIn(sentence.text, text)

    def test_docx_macrostructure_color_and_metadata(self):
        _, complaint, _ = self.make_complaint()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'complaint.docx'
            render_docx(complaint, path)
            document = Document(path)
            text = '\n'.join(p.text for p in document.paragraphs)
            self.assertEqual(text, complaint_plain_text(complaint))
            lines = [p.text for p in document.paragraphs]
            self.assertEqual(lines[2:4], list(COMPLAINT_TITLE))
            self.assertEqual([p.text for p in document.paragraphs if p.style.name == 'KSRF Heading'],
                             [FACTS_HEADING, REASONING_HEADING, REQUEST_HEADING, 'Приложения'])
            header = document.styles['KSRF Complaint Header']
            self.assertEqual(header.font.size.pt, 11.5)
            self.assertEqual(header.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertGreater(header.paragraph_format.left_indent, 0)
            highlighted = [run.text for p in document.paragraphs for run in p.runs
                           if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
            self.assertIn('[УКАЗАТЬ АДРЕС]', highlighted)
            self.assertIn('[УКАЗАТЬ ДАТУ]', highlighted)
            self.assertNotIn('Адрес: ', highlighted)
            for prop in (document.core_properties.title, document.core_properties.subject,
                         document.core_properties.author, document.core_properties.keywords):
                self.assertNotIn('synthetic-', prop)
                self.assertNotIn('draft', prop)
                self.assertNotIn('evidence', prop)
            for section in document.sections:
                combined = '\n'.join(p.text for p in section.header.paragraphs + section.footer.paragraphs)
                self.assertNotIn('ПРОВЕРИТЬ', combined)
                self.assertNotIn('РАБОЧИЙ', combined)

    def test_placeholder_detection_retains_strict_release_blockers(self):
        value = 'Адрес [АДРЕС РЕГИСТРАЦИИ]; [НЕ ПРЕДОСТАВЛЕНО: Факты]; [56]'
        self.assertEqual(find_unresolved_placeholders(value),
                         ['[АДРЕС РЕГИСТРАЦИИ]', '[НЕ ПРЕДОСТАВЛЕНО: Факты]'])

    def test_basis_prayer_numbered_requests_and_annex_order_preserve_provenance(self):
        basis = 'На основании условной нормы о полномочиях Суда заявитель обращается со следующими требованиями.'
        requests = ['1. Проверить условное положение в связи с указанным нормативным дефектом на соответствие Конституции.',
                    '2. Указать на пересмотр дела при наличии законных оснований в случае принятия постановления.']
        payload = {'matter_id': 'synthetic', 'draft_id': 'prayer', 'sections': [
            {'code': 'requested_remedy', 'heading': 'Внутренние требования', 'sentences': [
                {'text': text, 'role': 'requested_remedy', 'evidence_ids': ['synthetic-evidence']}
                for text in requests]},
            {'code': 'request_basis', 'heading': 'Внутренняя подводка', 'sentences': [
                {'text': basis, 'role': 'norm_text', 'evidence_ids': ['synthetic-norm']}]},
            {'code': 'authorities', 'heading': 'Источники', 'sentences': [
                {'text': 'Условный источник поддерживает предел проверяемого тезиса.',
                 'role': 'legal_holding', 'evidence_ids': ['synthetic-authority']}]},
            {'code': 'enclosures', 'heading': 'Опись', 'sentences': ['1. Копия условного акта.']},
            {'code': 'review_notes', 'heading': 'Анализ', 'sentences': ['Каков точный предел условного правила?']},
        ]}
        _, complaint, _ = prepare_working_draft(payload)
        before = complaint.to_dict()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'complaint.docx'
            render_docx(complaint, path)
            document = Document(path)
            paragraphs = document.paragraphs
            lines = [p.text for p in paragraphs]
            start = lines.index(REQUEST_HEADING)
            self.assertEqual(lines[start:start + 7], [REQUEST_HEADING, basis, PRAYER_MARKER,
                requests[0].replace('1. ', '1.\t', 1), requests[1].replace('2. ', '2.\t', 1),
                'Приложения', '1.\tКопия условного акта.'])
            self.assertEqual(lines.count(basis), 1)
            self.assertNotIn('Внутренняя подводка', lines)
            self.assertNotIn('Каков точный предел условного правила?', lines)
            prayer = paragraphs[start + 2]
            self.assertEqual(prayer.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertTrue(prayer.style.font.bold)
            self.assertEqual(prayer.style.font.size.pt, 12)
            self.assertEqual(prayer.style.paragraph_format.first_line_indent, 0)
            self.assertTrue(prayer.style.paragraph_format.keep_with_next)
            self.assertEqual([p.style.name for p in paragraphs[start + 3:start + 5]],
                             ['KSRF Numbered Item', 'KSRF Numbered Item'])
            self.assertEqual('\n'.join(lines), complaint_plain_text(complaint))
        self.assertEqual(complaint.to_dict(), before)
        authority = next(block for block in complaint_blocks(complaint)
                         if block.text.startswith('Условный источник'))
        self.assertEqual(authority.kind, 'source')

    def test_existing_prayer_and_selected_request_count_are_not_rewritten(self):
        texts = ['прошу:', '1. Сохранить единственное выбранное требование.']
        _, complaint, _ = prepare_working_draft({'matter_id': 'synthetic', 'draft_id': 'single',
            'sections': [{'code': 'requested_remedy', 'heading': 'Требования', 'sentences': texts},
                         {'code': 'rights_analysis', 'heading': 'Довод', 'sentences': [
                             'Дословная условная цитата: «Какая гарантия предусмотрена?»']} ]})
        before = complaint.to_dict()
        blocks = complaint_blocks(complaint)
        self.assertEqual([b.text for b in blocks if b.kind == 'prayer'], ['прошу:'])
        text = complaint_plain_text(complaint)
        self.assertIn('1.\tСохранить единственное выбранное требование.', text)
        self.assertNotIn('2.', text)
        self.assertIn('«Какая гарантия предусмотрена?»', text)
        self.assertEqual(complaint.to_dict(), before)

    def test_legacy_standalone_prayer_without_colon_is_preserved_without_duplicate(self):
        _, complaint, _ = prepare_working_draft({'matter_id': 'synthetic', 'draft_id': 'legacy',
            'sections': [{'code': 'requested_remedy', 'heading': 'Требования',
                          'sentences': ['ПРОШУ', '1. Прошу проверить условное положение.']}]})
        before = complaint.to_dict()
        blocks = complaint_blocks(complaint)
        self.assertEqual([b.text for b in blocks if b.kind == 'prayer'], ['ПРОШУ'])
        self.assertEqual([b.text for b in blocks if b.text.startswith('1.')],
                         ['1. Прошу проверить условное положение.'])
        self.assertEqual(complaint.to_dict(), before)

    def test_mixed_adverse_material_blocks_export_without_deleting_legal_prose(self):
        original, complaint, gaps = self.make_complaint()
        sections = list(complaint.sections)
        adverse = sections[4]
        mixed = adverse.sentences[0].text + ' Высокий риск: ПРОВЕРИТЬ: опоры.'
        sections[4] = replace(adverse, sentences=(replace(adverse.sentences[0], text=mixed),))
        complaint = replace(complaint, sections=tuple(sections))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'complaint.docx'
            with self.assertRaises(PresentationError):
                render_docx(complaint, path)
            self.assertFalse(path.exists())
        self.assertEqual(complaint.sections[4].sentences[0].text, mixed)
        self.assertIn(mixed, _review_markdown(complaint, gaps))
        self.assertIn(adverse.sentences[0].text, complaint_plain_text(original))

    def test_strict_companion_report_is_hash_bound_and_detects_tampering(self):
        from ksrf.filing.release import _declared_file_errors, _OPTIONAL_ARTIFACTS
        original, _, _ = self.make_complaint()
        relative, _, magic = _OPTIONAL_ARTIFACTS['review_markdown']
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder).resolve()
            path = root / relative
            artifact = render_review_markdown(original, path)
            self.assertIn('Риск: требуется проверить', path.read_text())
            self.assertNotIn('Суд установил', path.read_text())
            record = {**artifact.to_dict(), 'relative_path': relative}
            kwargs = dict(pack_root=root, expected_relative_path=relative,
                          expected_status='complete', label='review_markdown', expected_magic=magic)
            self.assertEqual(_declared_file_errors(record, **kwargs)[0], [])
            path.write_text(path.read_text() + '\nИзменено\n')
            self.assertTrue(any('hash_mismatch' in error for error in
                                _declared_file_errors(record, **kwargs)[0]))


if __name__ == '__main__':
    unittest.main()
