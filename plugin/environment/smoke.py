#!/usr/bin/env python3
"""Exercise real document tools on generated, unapproved synthetic materials."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile


class SmokeFailure(RuntimeError):
    """A concrete runtime or artifact contract was not satisfied."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise SmokeFailure(message)


def run_json(command: list[str], workspace: Path, allowed: tuple[int, ...] = (0,)) -> dict:
    completed = subprocess.run(
        command, cwd=workspace, capture_output=True, text=True, timeout=180,
    )
    require(completed.returncode in allowed,
            f"{Path(command[1]).name}: exit {completed.returncode}; {completed.stderr[-1200:]}")
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise SmokeFailure(f"{Path(command[1]).name}: invalid JSON output") from exc
    require(isinstance(payload, dict), "CLI result must be an object")
    return payload


def file_sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def exercise(skills: Path, workspace: Path, expect_unprivileged: bool) -> dict:
    from docx import Document
    from pypdf import PdfReader

    cycle = skills / "ksrf-complaint-cycle"
    cli = cycle / "scripts/ksrf.py"
    collect = cycle / "scripts/ksrf_autocollect.py"
    require(cli.is_file() and collect.is_file(), "Bundled runtime is missing")
    require(not workspace.resolve().is_relative_to(skills.resolve().parent),
            "Smoke workspace must be outside the installed package")
    if expect_unprivileged:
        require(hasattr(os, "geteuid") and os.geteuid() != 0, "Smoke must run as non-root")

    binaries = {}
    for name in ("soffice", "pdftotext", "pdftoppm", "tesseract"):
        binaries[name] = shutil.which(name)
        require(binaries[name] is not None, f"Required executable missing: {name}")
    languages = subprocess.run(
        [binaries["tesseract"], "--list-langs"], capture_output=True, text=True, timeout=30,
    )
    require(languages.returncode == 0, "Tesseract language probe failed")
    require({"rus", "eng"} <= set(languages.stdout.splitlines()), "OCR needs rus and eng")

    source_dir = workspace / "inputs"
    source_dir.mkdir()
    source = source_dir / "synthetic-material.docx"
    document = Document()
    document.add_heading("Синтетический материал для технической проверки", 0)
    for index in range(5):
        document.add_paragraph(
            f"Учебный фрагмент {index + 1}. Это вымышленный материал без персональных данных. "
            "Он проверяет чтение документа, сохранность русского текста и создание рабочего проекта. "
            "Судебные обстоятельства, применённая норма и процессуальные сроки не установлены."
        )
    document.save(source)
    source_hash = file_sha(source)
    matter = workspace / "matter"
    matter_id = "synthetic-environment-smoke"
    start = run_json([sys.executable, str(cli), "start", "--profile", "basic",
                      "--matter-id", matter_id, "--workspace", str(matter), "--json"], workspace)
    # The CLI accepts a human identifier and returns the canonical, scoped ID.
    matter_id = start.get("matter", {}).get("matter_id")
    require(isinstance(matter_id, str) and bool(matter_id), "CLI did not return a matter ID")
    intake = run_json([sys.executable, str(cli), "intake", "--workspace", str(matter),
                       "--input", str(source), "--json"], workspace)
    require(intake.get("state") == "registered" and len(intake.get("records", [])) == 1,
            "Synthetic source was not registered")
    require(intake.get("external_transmission_performed") is False, "Unexpected external transmission")

    extracted = run_json([sys.executable, str(collect), str(source), "--no-ocr"], workspace)
    require(extracted.get("document_count") == 1, "DOCX collection lost the source")
    collected = extracted["documents"][0]
    require(collected.get("sha256") == source_hash and collected.get("text_chars", 0) > 500,
            "DOCX extraction or source hash mismatch")

    render_input = workspace / "render-input.json"
    render_input.write_text(json.dumps({
        "schema_version": "1.0.0",
        "complaint": {
            "matter_id": matter_id, "draft_id": "synthetic-working-draft",
            "title": "Синтетический рабочий проект",
            "sections": [{"code": "facts", "heading": "Факты", "sentences": [{
                "text": "Синтетический пример проверяет экспорт документа. Обстоятельства требуют проверки.",
                "role": "fact", "support_status": "pending", "evidence_ids": [],
                "note": "Это технический пример; юридические выводы отсутствуют.",
            }]}],
        },
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    rendered = run_json([sys.executable, str(cli), "render", "draft", "--workspace", str(matter),
                         "--payload", str(render_input), "--json"], workspace, allowed=(0, 3))
    result = rendered.get("result", {})
    require(rendered.get("state") == "working_draft_created",
            "Working draft failed: " + json.dumps(result, ensure_ascii=False)[:1800])
    require(result.get("human_review") == "pending", "Human review state changed")
    for flag in ("filing_authority", "approval_authority", "release_eligible"):
        require(result.get(flag) is False, f"Unexpected authority: {flag}")
    require(result.get("gaps"), "Unverified source must retain explicit gaps")
    for artifact in result["artifacts"]:
        path = Path(artifact["path"])
        require(path.resolve().is_relative_to(workspace.resolve()), "Artifact escaped smoke workspace")
        require(path.is_file() and file_sha(path) == artifact["sha256"], "Artifact hash mismatch")
    pdf = Path(result["pdf"]["path"])
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    visible_text = " ".join(pdf_text.split())
    require("РАБОЧИЙ ПРОЕКТ" in visible_text and "Не для подписания и подачи" in visible_text,
            "PDF lost the working-draft notice")
    previews = result.get("preview_paths", [])
    require(len(previews) == len(PdfReader(pdf).pages) and len(previews) > 0,
            "Missing PDF page previews")
    pdf_extract = run_json([sys.executable, str(collect), str(pdf), "--no-ocr"], workspace)
    require(pdf_extract["documents"][0]["text_chars"] > 500, "Poppler PDF extraction failed")

    scan = source_dir / "synthetic-scan.png"
    shutil.copyfile(previews[0], scan)
    ocr = run_json([sys.executable, str(collect), str(scan)], workspace)
    require(ocr.get("document_count") == 1 and ocr["documents"][0]["text_chars"] > 100,
            "Synthetic image OCR returned insufficient text")
    verified = run_json([sys.executable, str(cli), "render", "draft-status",
                         "--workspace", str(matter), "--json"], workspace, allowed=(0, 3))
    require(verified.get("state") == "working_draft_created"
            and verified.get("result", {}).get("artifacts_revalidated") is True,
            "Saved working draft could not be revalidated")
    strict = run_json([sys.executable, str(cli), "render", "status",
                       "--workspace", str(matter), "--json"], workspace, allowed=(0, 3))
    require(strict.get("state") == "blocked"
            and strict.get("result", {}).get("reason_code") == "strict_render_missing",
            "Working draft was incorrectly treated as a strict release")
    require(file_sha(source) == source_hash, "Original source changed")
    return {
        "schema_version": "1.0", "status": "passed", "scope": "synthetic_software_smoke",
        "python": sys.version.split()[0], "uid": os.geteuid() if hasattr(os, "geteuid") else None,
        "checks": ["cli_start", "intake", "docx_extraction", "working_docx_pdf_export",
                   "artifact_hashes", "pdf_notice", "page_previews", "poppler_extraction",
                   "image_ocr_rus_eng", "draft_revalidation", "strict_release_still_blocked"],
        "page_count": len(previews), "source_sha256": source_hash,
        "human_review": "pending", "filing_authority": False, "release_eligible": False,
        "legal_quality_assessed": False, "workspace": str(workspace),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skills-root", type=Path,
                        default=Path(__file__).resolve().parents[1] / "skills")
    parser.add_argument("--workspace", type=Path,
                        help="New output directory outside the plugin; retained after the run.")
    parser.add_argument("--expect-unprivileged", action="store_true")
    args = parser.parse_args()
    temporary = None
    try:
        if args.workspace:
            workspace = args.workspace.resolve()
            require(not workspace.is_relative_to(args.skills_root.resolve().parent),
                    "Workspace must be outside the installed package")
            workspace.mkdir(parents=True, exist_ok=False)
        else:
            temporary = tempfile.TemporaryDirectory(prefix="ksrf-environment-smoke-")
            workspace = Path(temporary.name)
        result = exercise(args.skills_root.resolve(), workspace, args.expect_unprivileged)
        if args.workspace:
            (workspace / "smoke-report.json").write_text(
                json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (SmokeFailure, OSError, ValueError, ImportError, subprocess.TimeoutExpired) as exc:
        print(json.dumps({"status": "failed", "error": str(exc),
                          "scope": "synthetic_software_smoke"}, ensure_ascii=False), file=sys.stderr)
        return 1
    finally:
        if temporary is not None:
            temporary.cleanup()


if __name__ == "__main__":
    raise SystemExit(main())
